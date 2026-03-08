"""
Text extraction service supporting PDF, DOCX, and image files.
"""
import io
import os
from pathlib import Path
from typing import Optional


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file using pdfplumber (preferred) or PyPDF2."""
    text_parts = []
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n\n".join(text_parts)
    except Exception:
        pass

    # Fallback to PyPDF2
    try:
        import PyPDF2
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return "\n\n".join(text_parts)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {e}")


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {e}")


def extract_text_from_image(file_path: str) -> str:
    """Extract text from an image using pytesseract OCR."""
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
        return text.strip()
    except Exception as e:
        raise ValueError(f"Failed to extract text from image: {e}")


def extract_text(file_path: str, file_type: str) -> str:
    """
    Extract text from a file based on its type.

    Args:
        file_path: Path to the uploaded file.
        file_type: MIME type or extension (pdf, docx, png, jpg, etc.)

    Returns:
        Extracted text string.
    """
    file_type_lower = file_type.lower()

    if "pdf" in file_type_lower:
        return extract_text_from_pdf(file_path)
    elif "docx" in file_type_lower or "word" in file_type_lower or "openxmlformats" in file_type_lower:
        return extract_text_from_docx(file_path)
    elif any(ext in file_type_lower for ext in ["png", "jpg", "jpeg", "tiff", "bmp", "gif", "image"]):
        return extract_text_from_image(file_path)
    else:
        # Try to read as plain text
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            raise ValueError(f"Unsupported file type: {file_type}")
